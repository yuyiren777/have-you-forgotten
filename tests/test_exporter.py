import datetime
from types import SimpleNamespace

from openpyxl import load_workbook

from core.exporter import export_schedules, render_html


def _schedule(**overrides):
    values = {
        "id": 7,
        "title": "复习数学<第二章>",
        "description": "完成习题",
        "date": datetime.date(2026, 8, 3),
        "start_time": None,
        "end_time": None,
        "location": "自习室 A",
        "notes": '{"description":"带计算器"}',
        "repeat_rule": "weekly:1,3",
        "urgency": 2,
        "status": "pending",
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def test_excel_export_is_formatted_and_readable(tmp_path):
    target = tmp_path / "日程清单.xlsx"

    export_schedules([_schedule()], str(target), "xlsx")
    sheet = load_workbook(target)["日程"]

    assert [cell.value for cell in sheet[1]][:4] == ["标题", "日期", "开始时间", "结束时间"]
    assert sheet["C2"].value == "12:00（默认）"
    assert sheet["F2"].value == "带计算器"
    assert sheet["G2"].value == "紧急"
    assert sheet["I2"].value == "每周周一、周三"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref


def test_excel_keeps_undated_todo_without_fake_date(tmp_path):
    target = tmp_path / "无日期待办.xlsx"

    export_schedules([_schedule(date=None, start_time=None)], str(target), "xlsx")
    sheet = load_workbook(target)["日程"]

    assert sheet["B2"].value is None
    assert sheet["C2"].value is None


def test_html_export_is_standalone_printable_and_escaped(tmp_path):
    target = tmp_path / "日程清单.html"

    export_schedules([_schedule()], str(target), "html")
    content = target.read_text(encoding="utf-8")

    assert "<!doctype html>" in content
    assert "@media print" in content
    assert "复习数学&lt;第二章&gt;" in content
    assert "12:00（默认）" in content
    assert "带计算器" in content


def test_empty_html_still_explains_that_there_are_no_schedules():
    content = render_html([], datetime.datetime(2026, 8, 2, 12, 0))

    assert "共 0 条" in content
    assert "没有可导出的日程" in content
