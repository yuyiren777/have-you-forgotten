"""Local schedule export to familiar Excel and HTML files."""

from __future__ import annotations

import datetime
import os
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from utils.schedule_content import clean_optional_text, readable_notes


STATUS_NAMES = {
    "pending": "待处理",
    "reminded": "已提醒",
    "completed": "已完成",
    "expired": "已过期",
}
URGENCY_NAMES = {0: "普通", 1: "重要", 2: "紧急"}
WEEKDAY_NAMES = {1: "周一", 2: "周二", 3: "周三", 4: "周四", 5: "周五", 6: "周六", 7: "周日"}
HEADERS = [
    "标题", "日期", "开始时间", "结束时间", "地点", "备注",
    "紧急度", "状态", "重复安排",
]


def _repeat_name(repeat_rule) -> str:
    rule = clean_optional_text(repeat_rule)
    if rule == "daily":
        return "每天"
    if rule.startswith("weekly:"):
        days = [
            WEEKDAY_NAMES.get(int(day), "")
            for day in rule.split(":", 1)[1].split(",")
            if day.isdigit()
        ]
        return "每周" + "、".join(day for day in days if day) if days else ""
    if rule.startswith("monthly:"):
        day = rule.split(":", 1)[1]
        return f"每月{day}日" if day.isdigit() else ""
    return "不重复"


def _schedule_row(schedule) -> list:
    return [
        clean_optional_text(schedule.title),
        schedule.date,
        schedule.start_time.strftime("%H:%M") if schedule.start_time else (
            "12:00（默认）" if schedule.date else ""
        ),
        schedule.end_time.strftime("%H:%M") if schedule.end_time else "",
        clean_optional_text(schedule.location) or "未填写",
        readable_notes(schedule.notes) or clean_optional_text(schedule.description) or "未填写",
        URGENCY_NAMES.get(schedule.urgency, "普通"),
        STATUS_NAMES.get(schedule.status, schedule.status),
        _repeat_name(schedule.repeat_rule),
    ]


def write_xlsx(schedules, file_path: str) -> None:
    """Create an Excel workbook optimized for viewing and filtering."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "日程"
    sheet.append(HEADERS)
    for schedule in schedules:
        sheet.append(_schedule_row(schedule))

    header_fill = PatternFill("solid", fgColor="2B7A78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(bottom=Side(style="thin", color="D7DEE2"))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 26

    urgency_fills = {
        "重要": PatternFill("solid", fgColor="FFF1C7"),
        "紧急": PatternFill("solid", fgColor="FFD9D9"),
    }
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border
        if row[1].value:
            row[1].number_format = "yyyy-mm-dd"
        urgency_fill = urgency_fills.get(row[6].value)
        if urgency_fill:
            row[6].fill = urgency_fill
            row[6].font = Font(
                bold=True,
                color="A43232" if row[6].value == "紧急" else "815D00",
            )

    widths = [28, 14, 16, 12, 22, 42, 12, 12, 18]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(file_path)


def render_html(schedules, generated_at: datetime.datetime | None = None) -> str:
    """Create a standalone, printable schedule list."""
    generated_at = generated_at or datetime.datetime.now()
    rows = []
    for schedule in schedules:
        values = _schedule_row(schedule)
        cells = "".join(f"<td>{escape(str(value or ''))}</td>" for value in values)
        urgency_class = (
            " urgent" if schedule.urgency == 2
            else " important" if schedule.urgency == 1
            else ""
        )
        rows.append(f'<tr class="schedule-row{urgency_class}">{cells}</tr>')
    body = "".join(rows) or '<tr><td colspan="9" class="empty">没有可导出的日程</td></tr>'
    headers = "".join(f"<th>{escape(header)}</th>" for header in HEADERS)
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>AI备忘录日程清单</title>
<style>
* {{ box-sizing: border-box; }}
body {{ margin: 0; color: #18242e; background: #f5f7f8; font-family: "Microsoft YaHei UI", sans-serif; }}
main {{ max-width: 1400px; margin: 0 auto; padding: 32px; }}
h1 {{ margin: 0 0 6px; font-size: 28px; letter-spacing: 0; }}
.meta {{ margin: 0 0 24px; color: #5e6e78; }}
.table-wrap {{ overflow-x: auto; background: #fff; border: 1px solid #d7dee2; border-radius: 6px; }}
table {{ width: 100%; border-collapse: collapse; min-width: 1050px; }}
th {{ padding: 12px; color: #fff; background: #2b7a78; text-align: left; white-space: nowrap; }}
td {{ padding: 12px; border-bottom: 1px solid #e5eaed; vertical-align: top; white-space: pre-wrap; }}
tr:last-child td {{ border-bottom: 0; }}
tr.important td:first-child {{ border-left: 5px solid #e7ae24; }}
tr.urgent td:first-child {{ border-left: 5px solid #d64b4b; }}
.empty {{ padding: 40px; text-align: center; color: #677780; }}
@media print {{ body {{ background: #fff; }} main {{ max-width: none; padding: 0; }} .table-wrap {{ border: 0; }} }}
</style>
</head>
<body><main>
<h1>AI备忘录日程清单</h1>
<p class="meta">共 {len(schedules)} 条 · 导出时间：{generated_at:%Y-%m-%d %H:%M}</p>
<div class="table-wrap"><table><thead><tr>{headers}</tr></thead><tbody>{body}</tbody></table></div>
</main></body></html>"""


def export_schedules(schedules, file_path: str, export_format: str) -> None:
    """Write an export atomically so failures cannot leave a partial file."""
    path = Path(file_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    try:
        if export_format == "xlsx":
            write_xlsx(schedules, str(temporary))
        elif export_format == "html":
            temporary.write_text(render_html(schedules), encoding="utf-8")
        else:
            raise ValueError(f"不支持的导出格式：{export_format}")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()
